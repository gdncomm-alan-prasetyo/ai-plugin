#!/usr/bin/env python3
"""Render test-cases/*.md into the GDN test-case upload template.

Model writes case CONTENT as markdown files; this script owns the Excel FORMAT.
Keeping them separate means column order, the instruction sheet, and the enums
stay correct regardless of what the model emits.

Upload target creates new rows and never updates, so only `status: draft` cases
render. Re-rendering an uploaded case would duplicate it in the target system.

Usage:
    render_test_cases.py validate       --input test-cases/
    render_test_cases.py render         --input test-cases/ --output tc.xlsx \\
                                        --service iam-service --author diko.raditya
    render_test_cases.py mark-uploaded  --input test-cases/ [--jira KEY]
    render_test_cases.py retire         --input test-cases/ (--jira KEY | --case FILE)

Retirement always ends with the file deleted:
  - supersede (a new draft with supersedes: X) → mark-uploaded deletes X
  - standalone remove (behaviour gone, no replacement) → retire deletes the file
Both print the target rows to delete by hand first; the target appends rows and
this script cannot remove them.
"""

import argparse
import re
import shutil
import sys
from pathlib import Path

def require_openpyxl():
    """Import openpyxl on demand. Only `render` needs it.

    Kept lazy so importers that reuse the loader/validator/status writers -
    sync_test_tiger.py - work on a machine with no Excel dependency at all.
    """
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment
    except ImportError:
        sys.exit(
            "ERROR: openpyxl required.\n"
            "  python3 -m pip install --user openpyxl"
        )
    return load_workbook, Alignment

# Column order fixed by template header row. Do not reorder.
COLUMNS = [
    "test_case",        # A
    "test_summary",     # B
    "preconditions",    # C
    "step_action",      # D
    "expected_result",  # E
    "actual_result",    # F - execution-time, always blank
    "jira",             # G
    "author",           # H
    "test_type",        # I
    "importance",       # J
]

EXPECTED_HEADERS = [
    "Test Case", "Test Summary", "Preconditions", "Step Action",
    "Expected Result", "Actual Result", "JIRA", "Author",
    "Test Type", "Importance",
]

VALID_TEST_TYPES = {"NEW_FEATURE", "REGRESSION", "REGRESSION_SUPPORT", "INTEGRATION"}
VALID_IMPORTANCE = {"P0", "P1", "P2", "P3"}
VALID_STATUS = {"draft", "uploaded", "retired"}

# Template's "How to fill": empty preconditions are " - ", not blank.
NO_PRECONDITION = " - "

SECTIONS = {
    "Summary": "test_summary",
    "Preconditions": "preconditions",
    "Steps": "step_action",
    "Expected Result": "expected_result",
}

DEFAULT_TEMPLATE = Path(__file__).resolve().parents[1] / "assets" / "template_test_case_upload.xlsx"
DATA_SHEET = "service-name"
METADATA_SHEET = "Metadata"


def parse_frontmatter(text, filename):
    """Parse flat `key: value` frontmatter. Returns (dict, body). No nesting."""
    if not text.startswith("---"):
        raise ValueError(f"{filename}: missing '---' frontmatter block")
    parts = text.split("---", 2)
    if len(parts) < 3:
        raise ValueError(f"{filename}: unterminated frontmatter block")

    meta = {}
    for lineno, line in enumerate(parts[1].strip().splitlines(), start=2):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if ":" not in line:
            raise ValueError(f"{filename}:{lineno}: not a 'key: value' line: {line!r}")
        key, _, value = line.partition(":")
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
            value = value[1:-1]
        if value in ("null", "~", ""):
            value = None
        meta[key.strip()] = value
    return meta, parts[2]


def parse_body(body, filename):
    """Split body into the four `## ` sections."""
    found = {}
    current = None
    buffer = []
    for line in body.splitlines():
        heading = re.match(r"^##\s+(.+?)\s*$", line)
        if heading:
            if current:
                found[current] = "\n".join(buffer).strip()
            current = heading.group(1)
            buffer = []
        elif current:
            buffer.append(line)
    if current:
        found[current] = "\n".join(buffer).strip()

    unknown = set(found) - set(SECTIONS)
    if unknown:
        raise ValueError(f"{filename}: unknown section(s): {sorted(unknown)}")
    return found


def load_cases(input_dir):
    """Load every .md under input_dir. Returns (cases, errors)."""
    cases = []
    errors = []
    files = sorted(Path(input_dir).glob("*.md"))
    if not files:
        errors.append(f"no .md files found in {input_dir}")
        return cases, errors

    for path in files:
        try:
            meta, body = parse_frontmatter(path.read_text(encoding="utf-8"), path.name)
            sections = parse_body(body, path.name)
        except ValueError as exc:
            errors.append(str(exc))
            continue

        case = {"_file": path, "_name": path.name}
        case.update(meta)
        for heading, key in SECTIONS.items():
            case[key] = sections.get(heading, "").strip()
        cases.append(case)

    return cases, errors


def validate(cases):
    """Return list of human-readable problems. Empty means valid."""
    errors = []
    by_name = {c["_name"]: c for c in cases}
    seen_titles = {}

    for case in cases:
        name = case["_name"]

        for field in ("test_case", "jira", "test_type", "importance", "status"):
            if not str(case.get(field) or "").strip():
                errors.append(f"{name}: frontmatter '{field}' is required")

        for heading, key in SECTIONS.items():
            if key == "preconditions":
                continue  # empty is legal, renders as " - "
            if not case.get(key):
                errors.append(f"{name}: section '## {heading}' is required and must be non-empty")

        test_type = str(case.get("test_type") or "").strip()
        if test_type and test_type not in VALID_TEST_TYPES:
            errors.append(
                f"{name}: test_type {test_type!r} invalid; must be one of {sorted(VALID_TEST_TYPES)}"
            )

        importance = str(case.get("importance") or "").strip()
        if importance and importance not in VALID_IMPORTANCE:
            errors.append(
                f"{name}: importance {importance!r} invalid; must be one of {sorted(VALID_IMPORTANCE)}"
            )

        status = str(case.get("status") or "").strip()
        if status and status not in VALID_STATUS:
            errors.append(
                f"{name}: status {status!r} invalid; must be one of {sorted(VALID_STATUS)}"
            )

        tiger_id = str(case.get("tiger_id") or "").strip()
        if tiger_id and not tiger_id.isdigit():
            errors.append(
                f"{name}: tiger_id {tiger_id!r} invalid; must be the numeric Test Tiger case id"
            )

        if str(case.get("actual_result") or "").strip():
            errors.append(
                f"{name}: 'actual_result' must stay empty - QA fills it at execution time"
            )

        supersedes = case.get("supersedes")
        if supersedes:
            if supersedes not in by_name:
                errors.append(f"{name}: supersedes {supersedes!r} but no such file exists")
            elif supersedes == name:
                errors.append(f"{name}: supersedes itself")

        title = str(case.get("test_case") or "").strip()
        if title:
            if title in seen_titles:
                errors.append(
                    f"{name}: duplicate test_case title, also in {seen_titles[title]} - "
                    f"upload target keys on this name"
                )
            seen_titles[title] = name

    return errors


def retire_list(cases):
    """Uploaded cases superseded by a draft. Need manual deletion in target."""
    by_name = {c["_name"]: c for c in cases}
    retire = []
    for case in cases:
        if case.get("status") != "draft":
            continue
        target = case.get("supersedes")
        if not target:
            continue
        old = by_name.get(target)
        if old and old.get("status") in ("uploaded", "retired"):
            retire.append((old["_name"], old.get("test_case", ""), case["_name"]))
    return retire


def verify_template(ws):
    actual = [ws.cell(1, c).value for c in range(1, len(EXPECTED_HEADERS) + 1)]
    normalized = [str(v).strip() if v is not None else "" for v in actual]
    if normalized != EXPECTED_HEADERS:
        sys.exit(
            "ERROR: template header row does not match expected format.\n"
            f"  expected: {EXPECTED_HEADERS}\n"
            f"  found:    {normalized}\n"
            "Template may have been updated. Re-check column mapping before rendering."
        )


def render(drafts, template_path, output_path, service, version, author):
    load_workbook, Alignment = require_openpyxl()
    shutil.copyfile(template_path, output_path)
    wb = load_workbook(output_path)

    if DATA_SHEET not in wb.sheetnames:
        sys.exit(f"ERROR: template missing {DATA_SHEET!r} sheet (found: {wb.sheetnames})")

    ws = wb[DATA_SHEET]
    verify_template(ws)

    if METADATA_SHEET in wb.sheetnames:
        meta = wb[METADATA_SHEET]
        meta["A2"] = service
        meta["B2"] = version or ""

    wrap = Alignment(wrap_text=True, vertical="top")

    for offset, case in enumerate(drafts):
        row = 2 + offset  # data starts row 2, under header
        values = {
            "test_case": str(case["test_case"]).strip(),
            "test_summary": case["test_summary"],
            "preconditions": case["preconditions"] or NO_PRECONDITION,
            "step_action": case["step_action"],
            "expected_result": case["expected_result"],
            "actual_result": "",  # QA fills at execution
            "jira": str(case["jira"]).strip(),
            "author": str(case.get("author") or author).strip(),
            "test_type": str(case["test_type"]).strip(),
            "importance": str(case["importance"]).strip(),
        }
        for col, key in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row, column=col)
            cell.value = values[key]
            cell.alignment = wrap

    # Excel caps tab names at 31 chars, forbids []:*?/\
    safe = service
    for ch in "[]:*?/\\":
        safe = safe.replace(ch, "-")
    safe = safe[:31] or DATA_SHEET
    if safe != DATA_SHEET and safe not in wb.sheetnames:
        ws.title = safe

    wb.save(output_path)
    return ws.title


def set_status(path, new_status):
    """Rewrite a case file's status field in place."""
    text = path.read_text(encoding="utf-8")
    updated, count = re.subn(
        r"^status:.*$", f"status: {new_status}", text, count=1, flags=re.MULTILINE
    )
    if count == 0:
        raise ValueError(f"{path.name}: no 'status:' line to update")
    path.write_text(updated, encoding="utf-8")


def set_frontmatter_field(path, key, value):
    """Set a flat frontmatter key in place, inserting it when absent.

    Unlike set_status, this does not require the key to exist - optional
    fields like `tiger_id` are written on first sync.
    """
    text = path.read_text(encoding="utf-8")
    rendered = "null" if value is None else str(value)
    updated, count = re.subn(
        rf"^{re.escape(key)}:.*$", f"{key}: {rendered}", text, count=1, flags=re.MULTILINE
    )
    if count == 0:
        parts = text.split("---", 2)
        if len(parts) < 3:
            raise ValueError(f"{path.name}: no frontmatter block to update")
        updated = "---" + parts[1].rstrip("\n") + f"\n{key}: {rendered}\n" + "---" + parts[2]
    path.write_text(updated, encoding="utf-8")


def clear_supersedes(path):
    """Set a case file's supersedes field to null in place."""
    text = path.read_text(encoding="utf-8")
    updated, count = re.subn(
        r"^supersedes:.*$", "supersedes: null", text, count=1, flags=re.MULTILINE
    )
    if count:
        path.write_text(updated, encoding="utf-8")


def cmd_validate(args, cases):
    print(f"valid=true cases={len(cases)} drafts={sum(1 for c in cases if c.get('status') == 'draft')}")


def select_drafts(cases, jira_filter):
    """Drafts, optionally narrowed to one ticket. Shared by render and mark-uploaded."""
    drafts = [c for c in cases if c.get("status") == "draft"]
    if jira_filter:
        drafts = [c for c in drafts if str(c.get("jira") or "").strip() == jira_filter]
    return drafts


def cmd_render(args, cases):
    if not args.output:
        sys.exit("ERROR: --output required for render")
    if not args.service:
        sys.exit("ERROR: --service required for render")
    if not args.author:
        sys.exit("ERROR: --author required for render")

    drafts = select_drafts(cases, args.jira)
    if not drafts:
        if args.jira:
            print(f"no draft cases for {args.jira} - nothing to render")
        else:
            print("no draft cases - nothing to render")
            print("every case is already uploaded or retired")
        return

    skipped = sum(1 for c in cases if c.get("status") == "draft") - len(drafts)
    if skipped:
        print(f"note: {skipped} draft case(s) from other tickets not included (--jira {args.jira})")

    template_path = Path(args.template)
    if not template_path.is_file():
        sys.exit(f"ERROR: template not found: {template_path}")

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    sheet_name = render(drafts, template_path, output_path, args.service, args.version, args.author)

    print(f"written={output_path}")
    print(f"cases={len(drafts)} (drafts only)")
    print(f"sheet={sheet_name}")
    print("actual_result=blank (QA fills at execution)")

    retire = retire_list(cases)
    if retire:
        print()
        print(f"RETIRE MANUALLY IN TEST MANAGEMENT SYSTEM ({len(retire)}):")
        print("  upload target never updates, so these old rows stay live until deleted")
        for old_file, old_title, by_file in retire:
            print(f"  - {old_title}")
            print(f"      superseded by {by_file}")
    print()
    print("after uploading, run: mark-uploaded")


def cmd_mark_uploaded(args, cases):
    drafts = select_drafts(cases, args.jira)
    if not drafts:
        print(f"no draft cases to mark{f' for {args.jira}' if args.jira else ''}")
        return

    # Cases superseded by a now-uploaded draft. Their target rows were deleted
    # manually at render time; delete the local files now to keep the tree clean.
    # (old_file, _, by_file) — by_file is the superseding draft.
    retirements = [(old, by) for old, _, by in retire_list(cases)]
    by_name = {c["_name"]: c["_file"] for c in cases}

    for case in drafts:
        set_status(case["_file"], "uploaded")
    for old_name, by_name_key in retirements:
        by_name[old_name].unlink()
        # The supersede is consummated; clear the pointer so validate() does not
        # later fail on a supersedes: that names the now-deleted file.
        clear_supersedes(by_name[by_name_key])

    print(f"marked_uploaded={len(drafts)}")
    if retirements:
        print(f"deleted_superseded={len(retirements)}")
        for old_name, _ in sorted(retirements):
            print(f"  - {old_name}")


def cmd_retire(args, cases):
    """Standalone removal: behaviour deleted, no replacement case.

    Prints the target rows to delete by hand (for cases already uploaded),
    then deletes the local .md files.
    """
    if not args.jira and not args.case:
        sys.exit("ERROR: retire needs --jira <KEY> or --case <filename>")

    if args.case:
        selected = [c for c in cases if c["_name"] == args.case]
        if not selected:
            sys.exit(f"ERROR: no case file named {args.case!r} in {args.input}")
    else:
        selected = [c for c in cases if str(c.get("jira") or "").strip() == args.jira]
        if not selected:
            print(f"no cases for {args.jira} - nothing to retire")
            return

    # Uploaded (or already-retired) cases have a live row in the target system.
    uploaded = [c for c in selected if c.get("status") in ("uploaded", "retired")]
    draft = [c for c in selected if c.get("status") == "draft"]

    if uploaded:
        print(f"DELETE MANUALLY IN TEST MANAGEMENT SYSTEM ({len(uploaded)}):")
        print("  these rows are live in the target and this script cannot remove them")
        for c in uploaded:
            tiger_id = str(c.get("tiger_id") or "").strip()
            hint = f"  [Test Tiger case {tiger_id} - `testcases_cli.py delete-case {tiger_id}`]" if tiger_id else ""
            print(f"  - {c.get('test_case', c['_name'])}{hint}")
        print()

    for c in selected:
        c["_file"].unlink()

    print(f"deleted_local={len(selected)}")
    for c in sorted(selected, key=lambda x: x["_name"]):
        tag = "was uploaded - delete its target row above" if c.get("status") in ("uploaded", "retired") else "draft only - never uploaded, safe"
        print(f"  - {c['_name']} ({tag})")
    if draft and not uploaded:
        print("all draft-only: nothing to delete in the target system")


def main():
    parser = argparse.ArgumentParser(description="Render test-cases/*.md into GDN upload template")
    parser.add_argument("command", choices=["validate", "render", "mark-uploaded", "retire"])
    parser.add_argument("--input", required=True, help="Directory of test-case .md files")
    parser.add_argument("--output", help="Path for generated .xlsx (render only)")
    parser.add_argument("--service", help="Service name (render only)")
    parser.add_argument("--version", default="", help="Version stamp (render only)")
    parser.add_argument("--author", help="GDN username fallback (render only)")
    parser.add_argument("--jira", help="Only this ticket's cases (render, mark-uploaded, retire)")
    parser.add_argument("--case", help="A single case filename (retire only)")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="Override template path")
    args = parser.parse_args()

    input_dir = Path(args.input)
    if not input_dir.is_dir():
        sys.exit(f"ERROR: input directory not found: {input_dir}")

    cases, load_errors = load_cases(input_dir)
    errors = load_errors + (validate(cases) if cases else [])

    if errors:
        print(f"VALIDATION FAILED ({len(errors)} problem(s)):", file=sys.stderr)
        for err in errors:
            print(f"  - {err}", file=sys.stderr)
        sys.exit(1)

    {
        "validate": cmd_validate,
        "render": cmd_render,
        "mark-uploaded": cmd_mark_uploaded,
        "retire": cmd_retire,
    }[args.command](args, cases)


if __name__ == "__main__":
    main()
